"""أقفال ترقية محرّك جودة التقرير — Report Quality Engine Upgrade lock-tests.

> **العقد الحاكم (أمر العمل الرئيس).** كل عائلة عيبٍ رصدها تدقيق المالك التحريري
> على تقرير زبدة الفول السوداني/اليمن تصير **قاعدة عقد كاتب + قفل اختبار دائم**
> — لا تصحيحاً يدوياً لتقرير واحد. مدوّنة الإعادة الإنتاجية: `tools/canonical_yemen.py`
> (نفس شكل الإنتاج المخزَّن، تحمل كل العيوب: رمز HS خاطئ، أرقام قديمة، طلب Trends
> ضعيف، صفوف أسعار بلا وزن، HHI فوق رمز مُعلَّم).
>
> الموجات: 1 اتساق منطقي · 2 جودة البيانات · 3 المنافسة/التسعير · 4 العرض/البنية
> · 5 اللغة/النبرة · 6 عقد الحكم/خارطة الطريق.
"""
from __future__ import annotations

import importlib
import os
import sys
from unittest import mock

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tools.canonical_yemen import yemen_research_blob  # noqa: E402


# ════════════════════ الموجة ١ — الاتساق المنطقي ════════════════════

def test_w1_1_verdict_badge_matches_body_verdict():
    """1.1 — شارة الحكم (verdict_label) تطابق جملة الحكم في المتن/المختصر.

    مصدر أسبقية واحد (AI-أولاً) يستهلكه الشارة والمختصر والنبرة — لا تناقض
    «التوصية بالدخول» في الشارة مع «مراقبة» في المتن."""
    import silk_render as R
    from silk_narrative import verdict_ar
    dr = R.build_view(yemen_research_blob())["deep_research"]
    badge = dr["verdict_label"]
    # المتن/المختصر يشتقّ من نفس v_raw عبر verdict_ar.
    v_raw = ((dr["verdict"].get("ai") or {}).get("verdict")
             or dr["verdict"].get("verdict"))
    body = verdict_ar(v_raw)
    # كلاهما WATCH => «مراقبة السوق» / «مراقبة …» — لا go/دخول في أحدهما دون الآخر.
    assert "مراقبة" in badge and "مراقبة" in body
    assert "الدخول" not in badge  # الشارة لا تقول «التوصية بالدخول» بينما المتن مراقبة


def test_w1_2_hs_confirm_flags_peanut_butter_but_not_valid_matches():
    """1.2 — عقد التأكيد يُعلِّم «زبدة الفول السوداني»/040510 (الصفة المميّزة
    «فول سوداني» غائبة) ولا يُعلِّم التطابقات الصحيحة."""
    from silk_hs_confirm import confirm_hs, is_flagged
    bad = confirm_hs("زبدة الفول السوداني", "040510")
    assert bad["confirmed"] is False and is_flagged(bad)
    assert "فول" in bad["missing_terms"] and "سوداني" in bad["missing_terms"]
    # الصفة المميّزة لا يمكن أن تخسر أمام تطابق «زبدة» العاري.
    for name, code in [("تمور", "080410"), ("زبدة", "040510"),
                       ("olive oil", "150910")]:
        c = confirm_hs(name, code)
        assert c["confirmed"] is not False, (name, code, c)


def test_w1_2_hs_confirm_no_fabrication_on_unknown_code():
    """1.2 — رمزٌ خارج البذرة لا يُعلَّم False كاذبة؛ confirmed=None (غير مؤكَّد)."""
    from silk_hs_confirm import confirm_hs, is_flagged
    c = confirm_hs("زبدة الفول السوداني", "000000")  # خارج البذرة تماماً
    assert c["confirmed"] is None and not is_flagged(c)


def test_w1_3_flagged_hs_reframe_single_methodology_note_and_conf_cap():
    """1.3/4.1 — رمز مُعلَّم => (أ) ملاحظة منهجية **واحدة** «مؤشر سياقي»،
    (ب) سقف ثقة الحكم، (ج) الرمز مُتاح للمُصدِّرات عبر hs_flagged."""
    import silk_render as R
    from silk_hs_confirm import CONTEXTUAL_TAG
    dr = R.build_view(yemen_research_blob())["deep_research"]
    assert dr["hs_flagged"] is True
    # ثقة الحكم مسقوفة عند 0.5 (كانت 0.55 في المدوّنة).
    assert dr["verdict"]["confidence"] <= 0.5
    # ملاحظة «مؤشر سياقي» تظهر مرة واحدة بالضبط في الحدود (لا تكرار في كل قسم).
    hits = [l for l in dr["limits"] if CONTEXTUAL_TAG in l]
    assert len(hits) == 1, dr["limits"]


def test_w1_3_confidence_cap_never_raises_confidence():
    """1.3 — السقف لا يرفع ثقةً أدنى منه أبداً (عقد عدم الاختلاق)."""
    import silk_render as R
    blob = yemen_research_blob()
    blob["deep_research"]["verdict"]["confidence"] = 0.3
    blob["deep_research"]["verdict"]["ai"]["confidence"] = 0.3
    dr = R.build_view(blob)["deep_research"]
    assert dr["verdict"]["confidence"] == 0.3


def test_w1_3_unflagged_code_is_not_reframed():
    """1.3 — رمز مؤكَّد لا يُطأطئ ثقةً ولا يضيف ملاحظة سياقية (لا إيجابية كاذبة)."""
    import silk_render as R
    from silk_hs_confirm import CONTEXTUAL_TAG
    blob = yemen_research_blob()
    blob["product"] = "تمور"
    blob["hs_code"] = "080410"
    blob.pop("hs_confirmation", None)  # يُحسَب حياً => مؤكَّد
    dr = R.build_view(blob)["deep_research"]
    assert dr["hs_flagged"] is False
    assert not any(CONTEXTUAL_TAG in l for l in dr["limits"])


def test_w1_2_writer_prompt_reframes_flagged_hs_once():
    """1.3/4.1 — موجّه الكاتب يحمل تعليمة تأطير «مؤشر سياقي» + «مرة واحدة»
    حين يُمرَّر عقد تأكيد غير مؤكَّد."""
    import silk_ai_judge as J
    captured = {}

    def _fake_call(system, user, **kw):
        captured["user"] = user
        return "## 1. الخلاصة التنفيذية\nنص."

    with mock.patch.object(J, "available", return_value=True), \
         mock.patch.object(J, "_call", side_effect=_fake_call):
        J.deep_report({}, "ملخّص", {"verdict": "WATCH"},
                      "زبدة الفول السوداني", "اليمن", hs_code="040510",
                      hs_confirmation={"confirmed": False,
                                       "code_desc": "زبدة (Butter)",
                                       "missing_terms": ["فول", "سوداني"]})
    u = captured["user"]
    assert "مؤشر" in u and "سياقي" in u
    assert "مرة واحدة" in u  # تعليمة عدم التكرار


# ════════════════════ الموجة ٢ — جودة البيانات ════════════════════

# ── القاعدة العامة: التقادُم من المصدر لا النثر (قرار المالك) ──

def test_w2_1_fact_year_reads_structured_provenance_not_prose():
    """المصدر البنيوي: حقل صريح / وسم year=YYYY (البنك الدولي) / سنة retrieved_at."""
    import silk_staleness as S
    assert S.fact_year({"value": 1, "data_year": 2013}) == 2013
    assert S.fact_year({"value": 1, "note": "NY.GDP.PCAP.CD year=2013",
                        "retrieved_at": "2026-07-15"}) == 2013  # الوسم يسبق الجلب
    assert S.fact_year({"value": 1, "note": "دخل الفرد",
                        "retrieved_at": "2018-12-31"}) == 2018
    assert S.fact_year({"value": 1, "note": "x", "retrieved_at": ""}) is None


def test_w2_1_stale_fact_years_from_yemen_provenance():
    """قائمة الحقائق المتقادِمة تُشتَقّ من الحقائق (2013/2018)، لا من النثر."""
    from silk_staleness import stale_fact_years
    missions = yemen_research_blob()["deep_research"]["missions"]
    allf = [f for v in missions.values() for f in v["findings"]]
    assert stale_fact_years(allf) == {2013, 2018}


def test_w2_1_stale_fact_tagged_regardless_of_phrasing():
    """حقيقة متقادِمة (2013) تُوسَم بأيّ صياغة: «في 2013»، «2013م»، بلا قوسين."""
    import silk_render as R
    S = {2013, 2018}
    for s in ["في 2013 بلغ الدخل.", "عام 2013م نُشر المسح.",
              "الدخل 2013 كان منخفضاً.", "بيانات البنك الدولي 2018 تشير."]:
        out = R._tag_stale_years(s, S)
        assert R._STALE_TAG in out, s
    # «2013م» تبقى لاحقتها ملاصقة (لا تتدلّى «م»).
    assert "2013م — بيانات 2013" in R._tag_stale_years("عام 2013م.", S)


def test_w2_1_fresh_year_never_tagged():
    """رقم سنة حديث (2023+) لا يُوسَم أبداً (ليس في قائمة المتقادِمة)."""
    import silk_render as R
    out = R._tag_stale_years("واردات 2023 نحو 12 مليون دولار.", {2013, 2018})
    assert R._STALE_TAG not in out


def test_w2_1_hs_heading_2008_never_tagged_no_stale_fact_behind_it():
    """رمز/بند HS 2008 (للمحضرات) لا يُوسَم — رمزٌ لا سنةَ حقيقة، وليس في القائمة."""
    import silk_render as R
    out = R._tag_stale_years("العائلة الصحيحة البند 2008 للمحضرات.", {2013, 2018})
    assert R._STALE_TAG not in out


def test_w2_1_food_word_year_untagged_when_no_stale_fact():
    """«الطعام 2013» بلا حقيقة متقادِمة خلفه => بلا وسم (لا false-positive على
    كلمة تنتهي بـ«عام»؛ ولا حقيقة في القائمة)."""
    import silk_render as R
    out = R._tag_stale_years("استهلاك الطعام 2013 مرتفع.", set())
    assert R._STALE_TAG not in out


def test_w2_1_keyword_backstop_does_not_tag_longer_number_prefix():
    """مراجعة الشيفرة #3 — «عام 20135» لا يُوسَم (2013 بادئة رقمٍ أطول، حدّ يمينيّ)."""
    import silk_render as R
    out = R._tag_stale_years("نشر عام 20135 تقرير.", set())
    assert R._STALE_TAG not in out
    assert "20135" in out  # الرقم كما هو بلا كسر


def test_w2_1_era_suffix_does_not_swallow_adjacent_word():
    """مراجعة الشيفرة #4 — «2013مايو» تبقى «مايو» سليمة (لا ابتلاع أوّل حرفها)."""
    import silk_render as R
    out = R._tag_stale_years("عام 2013مايو حدث.", {2013})
    assert "مايو" in out and R._STALE_TAG in out
    # «2013م» المفردة عند حدّ كلمة تبقى لاحقتها ملاصقة.
    assert "2013م — بيانات 2013" in R._tag_stale_years("عام 2013م.", {2013})


def test_w2_1_vintage_from_structured_data_year_field_not_prose():
    """قرار المالك — الفِنتيج يُقرأ من الحقل البنيويّ `data_year` أولاً، لا من
    وسمٍ نصّيّ: كومتريد 2018 (جُلبت اليوم) => متقادِمة، 2024 => حديثة."""
    from silk_staleness import fact_year, is_stale_fact
    c18 = {"value": 9, "source": "UN Comtrade",
           "note": "إجمالي استيراد 2018, USD", "retrieved_at": "2026-07-15",
           "data_year": 2018}
    c24 = {"value": 9, "source": "UN Comtrade",
           "note": "إجمالي استيراد 2024, USD", "retrieved_at": "2026-07-15",
           "data_year": 2024}
    assert fact_year(c18) == 2018 and is_stale_fact(c18) is True
    assert fact_year(c24) == 2024 and is_stale_fact(c24) is False


def test_w2_1_observation_date_widened_non_dec31_and_fetch_stamp_excluded():
    """مراجعة الشيفرة #3 — احتياط retrieved_at يقبل أيّ تاريخ رصدٍ ماضٍ (لا
    نهاية السنة فقط)، ويستبعد طابع الجلب للسنة الجارية."""
    from silk_staleness import fact_year
    assert fact_year({"value": 1, "retrieved_at": "2018-06-30"}) == 2018  # غير نهاية السنة
    assert fact_year({"value": 1, "retrieved_at": "2013-12-31"}) == 2013
    import datetime
    cur = datetime.date.today().year
    assert fact_year({"value": 1, "retrieved_at": f"{cur}-07-15"}) is None  # طابع جلب


def test_w2_1_collectors_set_data_year_field_no_year_marker():
    """قرار المالك — الجامعون (كومتريد deep-research + المنافس) يضبطون الحقل
    البنيويّ `data_year` ولا يكتبون «year=» في الملاحظة."""
    import silk_llm_runtime as RT
    from silk_staleness import fact_year

    class _M:
        m49 = 887
        name_en = "Yemen"
    with mock.patch.object(RT, "comtrade_trade",
                           return_value=[{"partnerCode": "0", "primaryValue": 5_000_000}]), \
         mock.patch.object(RT, "primary_value", return_value=5_000_000):
        out = RT._tool_comtrade_imports({"years": [2018]},
                                        {"hs_code": "040510", "market": _M()})
    assert out[0].data_year == 2018
    assert "year=" not in out[0].note                 # لا وسم نصّيّ
    assert fact_year(out[0]) == 2018                  # يُقرأ من الحقل
    # المنافس أيضاً يضبط data_year.
    import silk_data_layer_v2 as V2
    cdp = V2._competitor_dp("372", 4_000_000, 10_000_000, hs_code="040510",
                            market_label=887, year=2018)
    assert cdp.data_year == 2018 and "year=" not in cdp.note


def test_w2_1_wb_collector_sets_data_year_and_no_year_marker():
    """قرار المالك — جامع البنك الدولي يضبط data_year ويُسقِط وسم «year=»."""
    import silk_data_layer as DL
    from silk_staleness import fact_year
    payload = [{"page": 1}, [{"value": 1106, "date": "2013"}]]
    with mock.patch.object(DL, "_cached_get", return_value=payload):
        dp = DL._world_bank_for_year("YEM", "NY.GDP.PCAP.CD", 2013)
    assert dp.value == 1106 and dp.data_year == 2013 and "year=" not in dp.note
    assert fact_year(dp) == 2013


def test_w2_1_client_surfaces_strip_residual_year_marker():
    """belt-and-suspenders (الدرس ٣٣) — مدوّنة مخزَّنة قديمة بوسم «year=» في
    ملاحظة تُنقّى من سطح العميل (md + docx خاليان من «year=»)."""
    import silk_render as R
    import silk_reports as RP
    blob = yemen_research_blob()
    # حقّن ملاحظة قديمة الطراز بوسم year= في حقيقة معروضة.
    blob["deep_research"]["missions"]["trade_flow"]["findings"][0]["note"] = \
        "إجمالي استيراد 2018 year=2018"
    view = R.build_view(blob)
    md = RP._md_deep_research(view, [])
    assert "year=" not in md
    import tempfile, os
    p = os.path.join(tempfile.mkdtemp(), "c.docx")
    RP.render_client_docx(view, p)
    from docx import Document
    doc = Document(p)
    txt = " ".join(par.text for par in doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for c in row.cells:
                txt += " " + c.text
    assert "year=" not in txt


def test_w2_1_yemen_report_tags_2013_2018_but_not_2008_2023():
    """التقرير المعروض لليمن: 2013/2018 موسومتان، 2008 (بند) و2023 (حديث) لا."""
    import silk_render as R
    txt = R.build_view(yemen_research_blob())["deep_research"]["report"]["text"]
    for yr in ("2013", "2018"):
        idx = txt.find(yr)
        assert R._STALE_TAG in txt[idx:idx + 45], yr
    # 2008 (بند HS) و2023 (حديث) بلا وسم ملاصق.
    for yr in ("2008", "2023"):
        idx = txt.find(yr)
        assert R._STALE_TAG not in txt[idx:idx + 45], yr


def test_w2_1_stale_tag_is_disclosure_only_no_number_change():
    """2.1 — الوسم إفصاح فقط: لا يغيّر أي رقم (عقد عدم الاختلاق)."""
    import silk_render as R
    out = R._tag_stale_years("دخل الفرد 1106 دولار في 2013.", {2013})
    assert "1106" in out  # الرقم كما هو


def test_w2_1_writer_choke_point_tags_stale_fact_for_writer():
    """نقطة الاختناق (silk_ai_judge._facts): الحقيقة المتقادِمة تُذيَّل بوسم
    الإفصاح في مدخلات الكاتب مهما كانت صياغته لاحقاً."""
    import silk_ai_judge as J

    class _DP:
        def __init__(s, value, note, ra, data_year=None):
            s.value, s.note, s.retrieved_at = value, note, ra
            s.source, s.confidence, s.data_year = "World Bank", 0.9, data_year

    class _Rep:
        agent_name = "economic"
        failed = False
        # الحقيقة المتقادِمة تحمل الحقل البنيويّ data_year (لا وسم نصّيّ)؛
        # الحديثة data_year للسنة الجارية.
        findings = [_DP(1106, "دخل الفرد", "2026-07-15", 2013),
                    _DP(12_000_000, "واردات", "2026-07-15", 2026)]
    facts = J._facts([_Rep()])
    # الحقيقة المتقادِمة (2013) موسومة؛ الحديثة لا.
    assert "الأحدث المتاح" in facts and "بيانات 2013" in facts
    assert facts.count("الأحدث المتاح") == 1


def test_w2_2_seasonality_gap_surfaces_closure_step_once():
    """2.2 — فجوة الموسمية تظهر في «ما لم يكتمل» مع خطوة الإغلاق، مرة واحدة."""
    import silk_render as R
    from silk_trends_agent import SEASONALITY_GAP_CLOSURE
    limits = R.build_view(yemen_research_blob())["deep_research"]["limits"]
    hits = [l for l in limits if l == SEASONALITY_GAP_CLOSURE]
    assert len(hits) == 1, limits
    assert "ميدان" in SEASONALITY_GAP_CLOSURE or "موزّع" in SEASONALITY_GAP_CLOSURE


def test_w2_3_broaden_if_weak_reports_both_terms_framed():
    """2.3 — صفة دقيقة شبه معدومة + فئة أعمّ قوية => توسيع آلي يعيد كليهما
    مؤطَّراً «الطلب على الفئة موجود؛ الصفة الدقيقة غير مبحوثة»."""
    import silk_trends_agent as T
    from silk_data_layer import DataPoint

    def _fake_interest(kw, geo=None, tf="today 12-m"):
        if "فوائد" in kw:            # الفئة الأعمّ
            return DataPoint(100.0, "Google Trends", 0.7, f"broad {kw}", "2026")
        return DataPoint(0.4, "Google Trends", 0.7, f"exact {kw}", "2026")

    def _fake_ctx(kw, geo=None, tf="today 12-m"):
        return {"related_top": [{"label": "فوائد زبدة الفول السوداني",
                                 "value": 100}],
                "related_rising": [], "topics_rising": [], "regions": [],
                "confidence": 0.6, "note": ""}

    with mock.patch.object(T, "trends_interest", side_effect=_fake_interest), \
         mock.patch.object(T, "trends_context", side_effect=_fake_ctx):
        weak = _fake_interest("زبدة الفول السوداني")
        broadened = T.broaden_if_weak("زبدة الفول السوداني", None,
                                      "today 12-m", weak)
    assert broadened is not None and broadened.value == 100.0
    assert "الفئة موجود" in broadened.note and "غير مبحوثة" in broadened.note


def test_w2_3_broaden_loop_capped_at_two_candidates():
    """مراجعة الشيفرة #2 — حلقة التوسيع محدودة (SILK_TRENDS_BROADEN_MAX=2):
    لا تستعلم pytrends لكل المرشّحين. مرشّحون ضعاف => نداءان كحدّ أقصى."""
    import silk_trends_agent as T
    from silk_data_layer import DataPoint
    calls = []

    def _fake_interest(kw, geo=None, tf="today 12-m"):
        calls.append(kw)
        return DataPoint(0.4, "Google Trends", 0.7, kw, "2026")  # كلها ضعيفة

    def _fake_ctx(kw, geo=None, tf="today 12-m"):
        return {"related_top": [{"label": f"t{i}", "value": 1} for i in range(6)],
                "related_rising": [], "topics_rising": [], "regions": [],
                "confidence": 0.6, "note": ""}

    with mock.patch.object(T, "trends_interest", side_effect=_fake_interest), \
         mock.patch.object(T, "trends_context", side_effect=_fake_ctx):
        weak = DataPoint(0.4, "Google Trends", 0.7, "دقيق", "2026")
        res = T.broaden_if_weak("دقيق", None, "today 12-m", weak)
    assert res is None                       # لا مرشّح أقوى
    assert len(calls) <= 2, calls            # سقف الاستعلامات


def test_w2_3_no_broaden_when_exact_term_is_strong():
    """2.3 — الصفة الدقيقة القوية لا تُوسَّع (لا اختلاق طلب زائد)."""
    import silk_trends_agent as T
    from silk_data_layer import DataPoint
    strong = DataPoint(60.0, "Google Trends", 0.7, "strong", "2026")
    assert T.broaden_if_weak("زبدة الفول السوداني", None, "today 12-m",
                             strong) is None


# ════════════════════ الموجة ٣ — المنافسة والتسعير ════════════════════

def test_w3_1_price_row_reason_classifier():
    """3.1 — مصنِّف سبب غياب السعر/كجم: وزن غير مذكور / وحدة غامضة / قابل للحساب."""
    import silk_render as R
    assert R._price_row_reason("علبة 5 دولار") == "وزن غير مذكور"
    assert R._price_row_reason("6.5 دولار/كجم") == ""
    assert R._price_row_reason("سعر 12 يورو للكيلو") == ""
    assert R._price_row_reason("عبوة كبيرة") == "وحدة غامضة"
    assert R._price_row_reason("") == "وحدة غامضة"
    assert R._price_row_reason("علبة 340 غرام بـ 5 دولار") == ""  # سعر+وزن


def test_w3_1_yemen_price_rows_carry_per_row_reason_and_single_unlock():
    """3.1 — صفوف أسعار اليمن تحمل سبباً لكل صفّ + سطر فتح واحد."""
    import silk_render as R
    dr = R.build_view(yemen_research_blob())["deep_research"]
    rows = dr["price_rows"]
    reasons = {r["reason"] for r in rows}
    assert "وزن غير مذكور" in reasons          # «علبة 5 دولار»
    assert "" in reasons                        # «6.5 دولار/كجم» قابل للحساب
    # سطر الفتح الوحيد مذكور مرة واحدة كبنية.
    assert "بطاقة منتج" in dr["price_unlock"] and "التكلفة/كجم" in dr["price_unlock"]


def test_w3_2_hhi_present_but_context_only_under_flagged_code():
    """3.2 — رمز مُعلَّم => التركّز حاضر لكنه سياقٌ فقط (لا إشارة تسجيل)، وثقة
    الحكم مسقوفة (لا يرفعها التركّز)."""
    import silk_render as R
    dr = R.build_view(yemen_research_blob())["deep_research"]
    assert dr["hs_flagged"] is True
    assert dr["concentration_context_only"] is True
    # HHI ما زال حاضراً في المتن/البعثات (لا يُحذَف — يُعاد تأطيره فقط).
    assert "HHI" in dr["report"]["text"] or "3100" in dr["report"]["text"]
    assert dr["verdict"]["confidence"] <= 0.5


def test_w3_2_unflagged_code_concentration_is_not_context_only():
    """3.2 — رمز مؤكَّد => التركّز إشارة عادية (context_only=False)."""
    import silk_render as R
    blob = yemen_research_blob()
    blob["product"] = "تمور"
    blob["hs_code"] = "080410"
    blob.pop("hs_confirmation", None)
    dr = R.build_view(blob)["deep_research"]
    assert dr["concentration_context_only"] is False


# ════════════════════ الموجة ٤ — العرض والبنية ════════════════════

def test_w4_1_hs_methodology_note_appears_at_most_once():
    """4.1 — تحذير رمز HS المنهجي يظهر مرة واحدة كحدّ أقصى في الحدود (لا تكرار)."""
    import silk_render as R
    from silk_hs_confirm import CONTEXTUAL_TAG
    limits = R.build_view(yemen_research_blob())["deep_research"]["limits"]
    assert sum(1 for l in limits if CONTEXTUAL_TAG in l) <= 1


def test_w4_1_writer_prompt_uses_short_backreference_for_hs_note():
    """4.1 — عقد الكاتب يوجّه لإحالة موجزة لا تكرار. المرساة صارت «قسم
    المنهجية» بدل «الملاحظة المنهجية» (بلاغ Nadec/اليمن #7): الأخيرة كانت
    تُطلِق `dangling_cross_reference` (FAIL) لأن الحارس لا يجد نصّها الحرفي؛
    «قسم المنهجية» مرساةٌ يقبلها الحارس فيُنهي الحجب — نفس مقصد الإحالة."""
    import silk_ai_judge as J
    captured = {}
    with mock.patch.object(J, "available", return_value=True), \
         mock.patch.object(J, "_call",
                           side_effect=lambda s, u, **k: captured.setdefault("u", u) or "## 1. خلاصة\nx"):
        J.deep_report({}, "م", {"verdict": "WATCH"}, "زبدة الفول السوداني",
                      "اليمن", hs_code="040510",
                      hs_confirmation={"confirmed": False, "code_desc": "زبدة",
                                       "missing_terms": ["فول"]})
    assert "انظر قسم المنهجية" in captured["u"]
    assert "انظر الملاحظة المنهجية" not in captured["u"]


def test_w1_3_writer_prompt_names_import_price_as_contextual_and_explains_contradiction():
    """1.3 (أمر التثبيت 2026-07-21) — تقرير الكويت الحيّ أظهر تناقض سعرٍ
    (تجزئة 0.67$/كجم < استيراد/جملة 6$/كجم) لأن عقد الكاتب السابق أطّر
    «حجم الاستيراد/CAGR/HHI/الحصص» فقط، بلا ذكر صريح لبند سعر الاستيراد —
    فالكاتب رصد التناقض نثراً بلا تفسيرٍ بنيوي. العقد الآن يسمّي «متوسط سعر
    الاستيراد/الجملة» صراحةً ويوجّه لتفسير التناقض بسبب فئة كومتريد مجاورة
    لا رقمٍ مُصلَح تخميناً."""
    import silk_ai_judge as J
    captured = {}
    with mock.patch.object(J, "available", return_value=True), \
         mock.patch.object(J, "_call",
                           side_effect=lambda s, u, **k: captured.setdefault("u", u) or "## 1. خلاصة\nx"):
        J.deep_report({}, "م", {"verdict": "WATCH"}, "زبدة الفول السوداني",
                      "الكويت", hs_code="040510",
                      hs_confirmation={"confirmed": False, "code_desc": "زبدة",
                                       "missing_terms": ["فول", "سوداني"]})
    u = captured["u"]
    assert "متوسط سعر الاستيراد" in u
    assert "ليس خطأً يُصحَّح تخميناً" in u


def test_w4_2_canonical_section_order_matches_target_sequence():
    """4.2 — الترتيب القانوني للأقسام يطابق التسلسل المطلوب في أمر العمل."""
    from silk_ai_judge import _REPORT_SECTIONS
    target = ["الخلاصة التنفيذية", "منهجية البحث ونطاقه",
              "نظرة عامة على السوق وحجمه", "ديناميكيات السوق",
              "تحليل المستهلك والطلب", "المشهد التنافسي",
              "التنظيم والوصول للسوق", "اللوجستيات وسلسلة الإمداد",
              "تقييم المخاطر", "التوصيات الاستراتيجية"]
    assert list(_REPORT_SECTIONS)[:len(target)] == target


def test_w4_3_style_contract_carries_length_budget():
    """4.3 — عقد الأسلوب يحمل ميزانية طول (~٣٠٪ أوجز) بقصّ التكرار لا الدليل."""
    import silk_style_contract as SC
    assert SC.TARGET_TIGHTEN_PCT == 30
    assert "٣٠" in SC.PROFESSIONAL_TONE_RULE or "30" in SC.PROFESSIONAL_TONE_RULE
    assert "الدليل" in SC.PROFESSIONAL_TONE_RULE  # لا يقصّ الدليل
    assert SC.PROFESSIONAL_TONE_RULE in SC.WRITER_STYLE_CONTRACT


# ════════════════════ الموجة ٥ — اللغة والنبرة ════════════════════

def test_w5_1_reviewer_flags_alarmist_tone_as_nonblocking_issue():
    """5.1 — المراجع يعلِّم النبرة التنبيهية مشكلةَ أسلوب (issues لا blocking)،
    بلا نداء مدفوع (فحص حتمي)."""
    import silk_ai_judge as J
    draft = ("## 1. الخلاصة التنفيذية\nيجب التوقف هنا فوراً فهذا يبطل كل "
             "الأرقام في سوق مضطربة وشحيحة البيانات.")
    # المراجع الـLLM يُحاكى None (لا مفتاح) => الفحص الحتمي وحده يعمل.
    with mock.patch.object(J, "available", return_value=True), \
         mock.patch.object(J, "_traced_call", return_value=None):
        rev = J.review_report(draft, {})
    assert rev is not None
    joined = " ".join(rev["issues"])
    assert "تنبيهية" in joined
    # ليست حاجبة (لا دورة تنقيح مدفوعة إضافية).
    assert not any("تنبيهية" in b for b in rev["blocking"])


def test_w5_1_measured_tone_draft_has_no_alarmist_issue():
    """5.1 — نص بنبرة مقيسة لا يُعلَّم (لا إيجابية كاذبة)."""
    import silk_ai_judge as J
    draft = ("## 1. الخلاصة التنفيذية\nينبغي التعامل مع هذه الأرقام كمؤشر "
             "سياقي لا كمقياس مباشر حتى تأكيد الرمز.")
    assert J._alarmist_issues(draft) == []


def test_w5_1_yemen_narrative_free_of_banned_phrases_and_rule_in_contract():
    """5.1 (طلب المالك) — السرد المعروض لمدوّنة اليمن خالٍ من العبارات الثلاث
    المحظورة **وأقرب صيغها** (مثل «يجب التوقف فوراً»)، ونصّ قاعدة النبرة موجود
    في ملف العقد."""
    import silk_render as R
    import silk_style_contract as SC
    txt = R.build_view(yemen_research_blob())["deep_research"]["report"]["text"]
    banned = ["يجب التوقف هنا فوراً", "يبطل كل الأرقام",
              "سوق مضطربة وشحيحة البيانات", "يجب التوقف فوراً",
              "تبطل كل الأرقام", "سوق مضطربة"]
    present = [p for p in banned if p in txt]
    assert present == [], present
    # كل صيغة محظورة مغطّاة بقائمة الإنفاذ (فحص المراجع الحتمي يمسكها).
    for p in banned:
        assert p in SC.ALARMIST_PHRASES, p
    # نصّ قاعدة النبرة + البديل المقيس حاضران في ملف العقد.
    assert SC.PROFESSIONAL_TONE_RULE in SC.WRITER_STYLE_CONTRACT
    assert "مؤشر سياقي لا كمقياس مباشر" in SC.MEASURED_TONE_HINT


def test_w5_2_sentence_length_guidance_present():
    """5.2 — إرشاد طول الجملة (تفضيل القِصار) في العقد، لا حدّ محارف صلب."""
    import silk_style_contract as SC
    assert SC.SENTENCE_MAX_WORDS == 25
    assert "القصير" in SC.PROFESSIONAL_TONE_RULE or "تُقسَم" in SC.PROFESSIONAL_TONE_RULE


# ════════════════════ الموجة ٦ — عقد الحكم وخارطة الطريق ════════════════════

def test_w6_1_watch_verdict_has_structured_flip_conditions():
    """6.1 — حكم «مراقبة» => شرطا قلب الحكم حقلان مهيكلان: بيانات تحت الرمز
    الصحيح (لأن الرمز مُعلَّم) + موزّع محلي مؤكَّد."""
    import silk_render as R
    dr = R.build_view(yemen_research_blob())["deep_research"]
    flips = dr["flip_conditions"]
    assert len(flips) == 2
    conds = " ".join(c["condition"] for c in flips)
    assert "الصحيح" in conds and "موزّع" in conds
    # كل شرط بخطوة إغلاق + حالة تحقّق (لا اختلاق: الموزّع غير محقَّق بلا جهات اتصال).
    for c in flips:
        assert c["closes_via"] and c["met"] is False


def test_w6_1_go_verdict_has_no_flip_conditions():
    """6.1 — حكم GO لا يحمل شرطي قلب (البنية مشروطة بمراقبة/مشروط فقط)."""
    import silk_render as R
    blob = yemen_research_blob()
    blob["deep_research"]["verdict"] = {"verdict": "GO", "confidence": 0.8,
                                        "ai": {"verdict": "GO", "reasoning": "x"}}
    # رمز مؤكَّد كي لا يُسقَف/يُعلَّم.
    blob["product"] = "تمور"
    blob["hs_code"] = "080410"
    blob.pop("hs_confirmation", None)
    dr = R.build_view(blob)["deep_research"]
    assert dr["flip_conditions"] == []


def test_w6_1_flip_conditions_render_in_markdown_export():
    """6.1 — شرطا قلب الحكم يظهران في تصدير Markdown تحت عنوانهما."""
    import silk_render as R
    import silk_reports as RP
    from silk_render import FLIP_CONDITIONS_HEADING
    view = R.build_view(yemen_research_blob())
    md = RP.render_markdown(view) if hasattr(RP, "render_markdown") else None
    if md is None:
        md = RP._md_deep_research(view, [])
    assert FLIP_CONDITIONS_HEADING in md
    assert "الصحيح" in md and "موزّع" in md


def test_w6_1_distributor_condition_met_when_confirmed_lead_exists():
    """6.1 — لا اختلاق: وجود موزّع بجهة اتصال مؤكَّدة => شرط الموزّع محقَّق."""
    import silk_render as R
    blob = yemen_research_blob()
    blob["deep_research"]["importer_leads"] = {
        "leads": [{"title": "موزّع عدن", "phone": "+967-1-000000"}],
        "path": "scraper"}
    dr = R.build_view(blob)["deep_research"]
    dist = [c for c in dr["flip_conditions"] if "موزّع" in c["condition"]][0]
    assert dist["met"] is True


def test_w6_1_filler_contact_does_not_satisfy_distributor_condition():
    """مراجعة الشيفرة #4 — رائد بهاتف/إيميل حشو («—»/«غير متاح») لا يُثبِت
    موزّعاً مؤكَّداً (عقد عدم الاختلاق)."""
    import silk_render as R
    assert R._real_contact("—") is False
    assert R._real_contact("") is False
    assert R._real_contact("غير متاح") is False
    # مراجعة الشيفرة #4 — عبارات الحشو المطوّلة تُرفَض أيضاً (رفض بنيوي).
    assert R._real_contact("غير متاح حالياً") is False
    assert R._real_contact("لا يوجد رقم") is False
    assert R._real_contact("+967-1-000000") is True
    assert R._real_contact("info@dist.ye") is True  # بريد
    blob = yemen_research_blob()
    blob["deep_research"]["importer_leads"] = {
        "leads": [{"title": "موزّع عدن", "phone": "غير متاح حالياً",
                   "email": "لا يوجد"}],
        "path": "scraper"}
    dr = R.build_view(blob)["deep_research"]
    dist = [c for c in dr["flip_conditions"] if "موزّع" in c["condition"]][0]
    assert dist["met"] is False


def test_w6_2_writer_prompt_caps_exec_summary_and_requires_flip_and_risks():
    """6.2 — عقد الكاتب يفرض سقف صفحتين للخلاصة التنفيذية + الحكم + شرطي
    القلب + ٣ أرقام + ٣ مخاطر."""
    import silk_ai_judge as J
    captured = {}
    with mock.patch.object(J, "available", return_value=True), \
         mock.patch.object(J, "_call",
                           side_effect=lambda s, u, **k: captured.setdefault("u", u) or "## 1. خلاصة\nx"):
        J.deep_report({}, "م", {"verdict": "WATCH"}, "زبدة الفول السوداني",
                      "اليمن", hs_code="040510")
    u = captured["u"]
    assert "شرطا قلب الحكم" in u
    assert "صفحتين" in u
    assert "ثلاثة مخاطر" in u and "ثلاثة أرقام" in u


# ════════════════════ بوّابة /research (1.2 — قبل الإنفاق) ════════════════════

def _client(**env):
    import api as api_mod
    for k, v in env.items():
        if v is None:
            os.environ.pop(k, None)
        else:
            os.environ[k] = v
    importlib.reload(api_mod)
    return api_mod


def test_w1_2_research_gate_422_on_unconfirmed_hs_before_spend():
    """1.2 — مع تفعيل الصمّام، رمزٌ غير مؤكَّد => 422 hs_confirmation_needed
    **قبل** أيّ حجز، وتأكيد المستخدم (hs_confirmed) يتجاوزها."""
    pytest.importorskip("fastapi")
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient
    api_mod = _client(SILK_HS_CONFIRM_GATE="1", SILK_API_KEY=None,
                      SILK_REQUIRE_HS6=None, SILK_WORLD_MARKETS=None)
    client = TestClient(api_mod.create_app())
    with mock.patch("requests.get",
                    side_effect=OSError("net blocked for offline test")):
        r = client.post("/research",
                        json={"product": "زبدة الفول السوداني",
                              "market": "Yemen", "hs_code": "040510",
                              "async_run": False, "persist": False})
    assert r.status_code == 422
    assert r.json()["detail"]["error"] == "hs_confirmation_needed"
    # تأكيد المستخدم يتجاوز البوّابة (لا يعود نفس الخطأ).
    with mock.patch("requests.get",
                    side_effect=OSError("net blocked for offline test")):
        r2 = client.post("/research",
                         json={"product": "زبدة الفول السوداني",
                               "market": "Yemen", "hs_code": "040510",
                               "hs_confirmed": True,
                               "async_run": False, "persist": False})
    assert not (r2.status_code == 422
                and r2.json().get("detail", {}).get("error")
                == "hs_confirmation_needed")


def test_w1_2_research_gate_on_by_default_blocks_unconfirmed_hs():
    """1.2 — **فشل-آمن**: بلا أيّ صمّام (افتراضُ الإنتاج) يُحجَب رمزٌ غير مؤكَّد.

    البلاغ الحيّ (2026-07-21): تشغيلةُ `/research` مدفوعة على «زبدة الفول
    السوداني» مضت على 040510 لأن البوّابة كانت خلف صمّامٍ مُطفأ في الإنتاج.
    القانون: لا إنفاق صامت على رمزٍ خاطئ. البوّابة الآن مفعّلة افتراضياً."""
    pytest.importorskip("fastapi")
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient
    api_mod = _client(SILK_HS_CONFIRM_GATE=None, SILK_API_KEY=None,
                      SILK_REQUIRE_HS6=None, SILK_WORLD_MARKETS=None)
    client = TestClient(api_mod.create_app())
    with mock.patch("requests.get",
                    side_effect=OSError("net blocked for offline test")):
        r = client.post("/research",
                        json={"product": "زبدة الفول السوداني",
                              "market": "Yemen", "hs_code": "040510",
                              "async_run": False, "persist": False})
    assert r.status_code == 422
    assert r.json()["detail"]["error"] == "hs_confirmation_needed"


def test_w1_2_research_gate_explicit_off_disables_block():
    """1.2 — إطفاءٌ صريح (`SILK_HS_CONFIRM_GATE=0`) يُعطّل الحجب (مهرب واعٍ)."""
    pytest.importorskip("fastapi")
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient
    api_mod = _client(SILK_HS_CONFIRM_GATE="0", SILK_API_KEY=None,
                      SILK_REQUIRE_HS6=None, SILK_WORLD_MARKETS=None)
    client = TestClient(api_mod.create_app())
    with mock.patch("requests.get",
                    side_effect=OSError("net blocked for offline test")):
        r = client.post("/research",
                        json={"product": "زبدة الفول السوداني",
                              "market": "Yemen", "hs_code": "040510",
                              "async_run": False, "persist": False})
    assert not (r.status_code == 422
                and r.json().get("detail", {}).get("error")
                == "hs_confirmation_needed")


# ══════ الموجة ٢ (2026-07-21، أمر التثبيت): نقطة اختناق مشتركة — البوّابة
# تعمل على /analyze و/research معاً بلا صمّامٍ إضافي (نفس `preflight_block`
# المستدعى من كلا المعالجَين، silk_hs_confirm.py). إصلاحٌ سابقٌ اقتصر على
# /research وحده فأنتج تقرير كويت/زبدة الفول السوداني الحيّ — «إصلاحٌ على
# مسارٍ واحد نصفُ إصلاح». ══════

def test_w2_hs_gate_blocks_on_both_analyze_and_research_by_default():
    """الموجة ٢ — كلا المسارين (/analyze و/research) يحجبان رمزاً غير
    مؤكَّد **بلا أيّ صمّام مضبوط** (افتراض الإنتاج)؛ التقاعس عن أحدهما هو
    بالضبط الحادثة الحيّة (تقرير الكويت/زبدة الفول السوداني، 2026-07-21)."""
    pytest.importorskip("fastapi")
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient
    api_mod = _client(SILK_HS_CONFIRM_GATE=None, SILK_API_KEY=None,
                      SILK_REQUIRE_HS6=None, SILK_WORLD_MARKETS=None)
    client = TestClient(api_mod.create_app())
    with mock.patch("requests.get",
                    side_effect=OSError("net blocked for offline test")):
        r_research = client.post(
            "/research", json={"product": "زبدة الفول السوداني",
                               "market": "Kuwait", "hs_code": "040510",
                               "async_run": False, "persist": False})
        r_analyze = client.post(
            "/analyze", json={"product": "زبدة الفول السوداني",
                              "hs_code": "040510", "persist": False})
    for label, r in (("research", r_research), ("analyze", r_analyze)):
        assert r.status_code == 422, f"{label}: {r.status_code} {r.text}"
        assert r.json()["detail"]["error"] == "hs_confirmation_needed", label


def test_w2_hs_gate_choke_point_is_shared_not_duplicated():
    """الموجة ٢ — بنيوياً: كلّ معالجٍ ينفق يستدعي نفس دالة الاختناق
    `silk_hs_confirm.preflight_block` — لا منطق مكرَّر في كل معالج (لو
    اختفى الاستدعاء من أحدهم بصمت، هذا الحارس يفشل قبل أيّ بلاغ حيّ).

    اللائحة ٦٥: المعالجاتُ تستدعيها اليوم عبر الغلاف `preflight_resolve`
    (البوّابةُ نفسُها + قياسُ السمة الرقمية). يُحتسَب الغلافُ **بشرط** أن
    يُثبَت بنيوياً أنه يستدعي `preflight_block` لا يستبدلها — وإلا لكان
    بوّابةً موازيةً تُفرِغ الحارس من معناه."""
    import inspect
    import silk_hs_confirm
    src = _read_api_source()
    calls = src.count("preflight_block(") + src.count("preflight_resolve(")
    assert calls >= 2, (
        "نقطة الاختناق المشتركة يجب أن تُستدعى من كلا "
        "معالجَي /analyze و/research")
    assert "preflight_block(" in inspect.getsource(
        silk_hs_confirm.preflight_resolve), (
        "preflight_resolve لا يستدعي preflight_block — بوّابةٌ موازية")


def _read_api_source() -> str:
    import os as _os
    root = _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))
    with open(_os.path.join(root, "api.py"), encoding="utf-8") as f:
        return f.read()
